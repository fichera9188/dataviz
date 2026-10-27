"""Inventory selected research products and validate the local source files."""
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
RAW = ROOT / 'source-data'
products = [
    ('D01', 'WHO total alcohol consumption per adult', 'WHO', '2000-2024 (verified live API extract)', ['D01_WHO_total_consumption_raw.json','D01_WHO_total_consumption_countries.csv'], 'https://ghoapi.azureedge.net/api/SA_0000001688', ['country','year','litres_pure_alcohol_per_person_15_plus'], 'CSV available; geography still needed'),
    ('D02', 'WHO heavy episodic drinking by sex', 'WHO', '2000-2020 (verified live API extract)', ['D02_WHO_heavy_episodic_drinking_raw.json','D02_WHO_heavy_episodic_drinking_countries.csv'], 'https://ghoapi.azureedge.net/api/SA_0000001739', ['country','year','sex','prevalence_percent'], 'CSV available; explain definition and country-level distribution'),
    ('D03', 'National Drug Strategy Household Survey alcohol', 'AIHW', 'Measure-specific; national 2001-2022/23 and state commonly 2007-2022/23', ['D03_NDSHS_alcohol.xlsx','D03_NDSHS_states.xlsx'], 'https://www.aihw.gov.au/reports/illicit-use-of-drugs/national-drug-strategy-household-survey-2022-2023/data', ['age','gender','year','drinking_status','risk','state'], 'Reshape selected tables; preserve reliability flags'),
    ('D04', 'National Health Survey alcohol by age and sex', 'ABS', '2022', ['D04_ABS_NHS_alcohol_age_sex.xlsx'], 'https://www.abs.gov.au/statistics/health/health-conditions-and-risks/national-health-survey/2022/NHSDC07.xlsx', ['age','sex','guideline_exceedance','estimates','proportions'], 'Reshape Table 7; adult conclusions use ages 18+'),
    ('D05', 'Secondary Students Alcohol and Drug survey', 'Australian Government Department of Health / Cancer Council Victoria and partners', '2022/23 plus published historical comparisons', ['D05_ASSAD_2022_2023.pdf'], 'https://www.health.gov.au/resources/publications/australian-secondary-school-students-use-of-alcohol-and-other-substances-2022-2023?language=en', ['age_12_to_17','sex_or_gender','drinking_period','supply','location'], 'PDF aggregate table extraction and verification required'),
    ('D06', 'Alcohol available for consumption in Australia', 'AIHW', 'Total pure alcohol 1960/61-2023/24; some beverage series differ', ['D06_alcohol_availability_2023_2024.xlsx'], 'https://www.aihw.gov.au/getmedia/103e7c27-595e-418f-8260-685c3c2dc3c7/PHE331_AACA_Data-tables-2023-24.xlsx', ['financial_year','beverage','pure_alcohol_volume','per_capita_availability'], 'Reshape; preserve unavailable cider history'),
    ('D07', 'Alcohol-induced deaths', 'ABS', '2015-2024; selected tables start 2020; recent years provisional/revised', ['D07_ABS_alcohol_induced_deaths_2024.xlsx','D07_ABS_revision_release_2026.xlsx'], 'https://www.abs.gov.au/statistics/health/causes-death/causes-death-australia/latest-release', ['year','sex','age','deaths','death_rate','acute_or_chronic'], 'Use alcohol-specific cube and retain revision status; May 2026 revision workbook inspected'),
    ('D08', 'Alcohol and other drug treatment services', 'AIHW', '2024/25 latest; selected historical state series 2015/16-2024/25', ['D08_treatment_clients_2024_2025.xlsx','D08_treatment_states_2024_2025.xlsx'], 'https://www.aihw.gov.au/reports/alcohol-other-drug-treatment-services/alcohol-other-drug-treatment-services-australia/data', ['state','age','sex','principal_drug','clients','closed_episodes'], 'Filter alcohol; clients and episodes are separate measures; no inferred joint age-sex data'),
    ('D09', 'Alcohol-attributable burden of disease', 'AIHW', '2024 (compact extract)', ['D09_alcohol_attributable_burden_2024.xlsx'], 'https://www.aihw.gov.au/getmedia/f6366c20-5686-454e-a3a9-3bbba5db789a/AIHW-PHE221_ABDS-tables_DEC25_1.xlsx', ['linked_disease','sex','attributable_healthy_life_years_lost','attributable_percent'], 'Exclude overlapping parent/component rows; include the explicit residual and rank positive counts'),
    ('D10', 'Alcohol-related hospitalisations', 'AIHW NHMD', '2015/16-2023/24; age-sex snapshot 2023/24', ['D10_alcohol_hospitalisations_2023_2024.xlsx'], 'https://www.aihw.gov.au/getmedia/f774a123-a571-4938-a7b7-f7957d0dc6ce/AIHW-PHE221_NHMD-tables_DEC25_1.xlsx', ['age','sex','year','hospitalisations','same_day_or_overnight','remoteness'], 'No state breakdown in extract; use actual NHMD3 joint age-sex table'),
]
catalog = []
for product_id, title, provider, coverage, names, source, fields, readiness in products:
    files = []
    for name in names:
        path = RAW / name
        if not path.exists():
            raise FileNotFoundError(path)
        data = path.read_bytes()
        record = {'path': str(path), 'bytes': len(data), 'sha256': hashlib.sha256(data).hexdigest()}
        if path.suffix == '.xlsx':
            if not data.startswith(b'PK'):
                raise ValueError(f'Invalid workbook response: {name}')
            book = load_workbook(path, read_only=True, data_only=True)
            record['worksheets'] = book.sheetnames
            book.close()
        elif path.suffix == '.pdf' and not data.startswith(b'%PDF'):
            raise ValueError(f'Invalid PDF response: {name}')
        elif path.suffix == '.json':
            json.loads(data)
        files.append(record)
    catalog.append({'id': product_id, 'title': title, 'provider': provider, 'coverage': coverage, 'source_url': source, 'fields': fields, 'preparation_status': readiness, 'local_files': files})
assert len(catalog) == 10
(ROOT / 'DATASET_CATALOG.json').write_text(json.dumps({'research_date': '2026-10-13', 'scope': 'Global context and Australian case study', 'counting_note': 'Ten data products; WHO indicators share one information system; national and state NDSHS workbooks count together; treatment workbooks count together.', 'datasets': catalog}, indent=2), encoding='utf-8')
print(json.dumps({'validated_products': len(catalog), 'validated_files': sum(len(p['local_files']) for p in catalog), 'all_local_sources_present': True, 'catalog': str(ROOT / 'DATASET_CATALOG.json')}, indent=2))
